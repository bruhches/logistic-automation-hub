from pathlib import Path
from datetime import datetime, date, time
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

class SpreadsheetProcessor:
    REQUIRED=['Placa Veículo','Nº Documento / Carga','Motorista','Data Entrada','Hora Entrada','Tipo Operação','Status']
    ID_NAMES={'id','id registro','registro','codigo','código'}
    def detect_header(self,path,sheet_name=0,scan_rows=20):
        raw=pd.read_excel(path,sheet_name=sheet_name,header=None,nrows=scan_rows,engine='openpyxl'); best=(0,-1)
        for idx,row in raw.iterrows():
            vals=[str(v).strip() for v in row.tolist() if pd.notna(v) and str(v).strip()]
            score=len(vals)+len(set(vals))*.15
            if len(vals)>=2 and score>best[1]: best=(idx,score)
        return int(best[0])
    def analyze(self,path:Path,rules=None):
        rules=rules or {}
        def rule(name, default):
            return rules.get(name, default)
        xls=pd.ExcelFile(path,engine='openpyxl'); sheet=xls.sheet_names[0]; h=self.detect_header(path,sheet)
        df=pd.read_excel(path,sheet_name=sheet,header=h,engine='openpyxl').dropna(how='all').dropna(axis=1,how='all'); df.columns=[str(c).strip() for c in df.columns]
        findings=[]
        def add(i,kind,severity,message,action='Revisar',related=None):
            fid=f'F{len(findings)+1:03d}'; hub_id=f'HUB-{int(i)+1:06d}'; item={'id':fid,'hub_id':hub_id,'row_index':int(i),'excel_row':h+2+int(i),'type':kind,'severity':severity,'message':message,'action':action,'decision':'Remover' if action=='Remover cópia' else 'Revisar','related_row':(h+2+int(related)) if related is not None else None,'related_index':int(related) if related is not None else None}
            findings.append(item)
        er=rule('exact_duplicate',{'enabled':True,'severity':'Crítico','default_decision':'Remover'})
        od=rule('operational_duplicate',{'enabled':True,'severity':'Atenção','default_decision':'Revisar','fields':[]})
        exact=df.duplicated(keep=False) if er.get('enabled',True) else pd.Series(False,index=df.index)
        configured=[c for c in od.get('fields',[]) if c in df.columns]
        opcols=configured or [c for c in df.columns if c.lower() not in self.ID_NAMES]
        op=df.duplicated(subset=opcols,keep=False) if od.get('enabled',True) and opcols else pd.Series(False,index=df.index)
        seen={}
        for i,row in df[exact].iterrows():
            key=tuple(self._key(v) for v in row.tolist())
            if key in seen:add(i,'Duplicidade exata',er.get('severity','Crítico'),f'Registro idêntico à linha {h+2+seen[key]}.','Remover cópia' if er.get('default_decision','Remover')=='Remover' else 'Revisar',seen[key])
            else:seen[key]=i
        seen={}
        for i,row in df[op].iterrows():
            key=tuple(self._key(row[c]) for c in opcols)
            if key in seen and not exact.loc[i]:add(i,'Possível duplicidade operacional',od.get('severity','Atenção'),f'Mesma operação da linha {h+2+seen[key]} pelos campos configurados: {', '.join(opcols)}. Exige revisão humana.','Revisar',seen[key])
            elif key not in seen:seen[key]=i
        rr=rule('required_fields',{'enabled':True,'severity':'Crítico','fields':self.REQUIRED})
        if rr.get('enabled',True):
            for c in rr.get('fields',self.REQUIRED):
                if c in df.columns:
                    for i in df.index[df[c].isna()|(df[c].astype(str).str.strip()=='')]:add(i,'Campo obrigatório',rr.get('severity','Crítico'),f'Campo obrigatório “{c}” não informado.')
        dr=rule('datetime_order',{'enabled':True,'severity':'Crítico'})
        if dr.get('enabled',True) and all(c in df.columns for c in ['Data Entrada','Hora Entrada','Data Saída','Hora Saída']):
            for i,r in df.iterrows():
                ent=self._combine(r['Data Entrada'],r['Hora Entrada']); sai=self._combine(r['Data Saída'],r['Hora Saída'])
                if ent and sai and sai<ent:add(i,'Data/Horário',dr.get('severity','Crítico'),'Data/hora de saída é anterior à entrada.')
        cr=rule('completed_without_exit',{'enabled':True,'severity':'Atenção','statuses':['concluído','concluido','finalizado','finalizada']})
        ar=rule('allowed_status',{'enabled':False,'severity':'Atenção','values':[]})
        if 'Status' in df.columns:
            allowed={str(x).strip().lower() for x in ar.get('values',[])}
            concluded={str(x).strip().lower() for x in cr.get('statuses',[])}
            for i,r in df.iterrows():
                st=str(r.get('Status','')).strip().lower()
                if cr.get('enabled',True) and st in concluded and (pd.isna(r.get('Data Saída')) or pd.isna(r.get('Hora Saída'))):add(i,'Status',cr.get('severity','Atenção'),'Registro concluído sem data/hora de saída.')
                if ar.get('enabled',False) and st and st not in allowed:add(i,'Status fora do padrão',ar.get('severity','Atenção'),f'Status “{r.get("Status")}” não está na lista configurada.')
        affected=len(set(x['excel_row'] for x in findings)); counts={s:sum(x['severity']==s for x in findings) for s in ['Crítico','Atenção','Informativo']}
        summary={'rows':len(df),'columns':len(df.columns),'critical':counts['Crítico'],'warning':counts['Atenção'],'info':counts['Informativo'],'affected_records':affected,'clean_records':max(len(df)-affected,0),'header_row':h+1}
        issues=[f'{len(findings)} ocorrência(s) em {affected} registro(s). Decisões de remoção/manutenção podem ser feitas antes do processamento.']
        if h>0:issues.append(f'Cabeçalho detectado na linha {h+1}; linhas anteriores serão preservadas.')
        preview=df.head(8)
        return {'sheet':sheet,'header_idx':h,'df':df,'findings':findings,'summary':summary,'issues':issues,'preview':{'columns':list(df.columns),'rows':[{c:self._json(v) for c,v in r.items()} for _,r in preview.iterrows()]},'columns':list(df.columns)}
    def comparison(self,analysis,finding_id):
        f=next((x for x in analysis['findings'] if x['id']==finding_id),None)
        if not f:return None
        df=analysis['df']; a=df.loc[f['row_index']]; b=df.loc[f['related_index']] if f['related_index'] is not None else None
        rows=[]
        for c in df.columns: rows.append({'field':c,'current':self._json(a[c]),'related':self._json(b[c]) if b is not None else ''})
        return {'finding':{k:v for k,v in f.items() if k not in {'row_index','related_index'}},'fields':rows}
    def process(self,source,output,analysis,decisions,process_code=None,processed_at=None):
        df=analysis['df'].copy(); remove=set()
        for f in analysis['findings']:
            decision=decisions.get(f['id'],f.get('decision','Revisar')); f['decision']=decision
            if decision=='Remover':remove.add(f['row_index'])
        kept_indices=[int(i) for i in df.index if int(i) not in remove]
        cleaned=df.loc[kept_indices].copy()
        final_rows={idx: analysis['header_idx']+2+pos for pos,idx in enumerate(kept_indices)}
        for f in analysis['findings']:
            f['final_excel_row']=final_rows.get(int(f['row_index']))
        wb=load_workbook(source); ws=wb[analysis['sheet']]; hr=analysis['header_idx']+1
        if ws.max_row>hr:ws.delete_rows(hr+1,ws.max_row-hr)
        for row in cleaned.itertuples(index=False,name=None):ws.append([self._excel(v) for v in row])
        fill=PatternFill('solid',fgColor='203040'); font=Font(color='FFFFFF',bold=True)
        for c in ws[hr]:c.fill=fill;c.font=font;c.alignment=Alignment(horizontal='center')
        for i,col in enumerate(cleaned.columns,1):
            vals=[str(col)]+[str(self._excel(v) or '') for v in cleaned.iloc[:,i-1].head(100)];ws.column_dimensions[get_column_letter(i)].width=min(max(max(map(len,vals))+2,10),35)
        ws.freeze_panes=f'A{hr+1}';ws.auto_filter.ref=f'A{hr}:{get_column_letter(len(cleaned.columns))}{hr+len(cleaned)}'
        if 'ANALISE_AUTOMACAO' in wb.sheetnames:del wb['ANALISE_AUTOMACAO']
        au=wb.create_sheet('ANALISE_AUTOMACAO'); headers=['ID Processamento','ID Hub','ID ocorrência','Linha original','Linha final','Linha relacionada','Tipo','Severidade','Problema','Decisão'];au.append(headers)
        for f in analysis['findings']:au.append([process_code or '',f['hub_id'],f['id'],f['excel_row'],f.get('final_excel_row'),f['related_row'],f['type'],f['severity'],f['message'],f.get('decision','Revisar')])
        for c in au[1]:c.fill=fill;c.font=font;c.alignment=Alignment(horizontal='center')
        for i,w in enumerate([18,16,14,16,14,18,34,14,75,18],1):au.column_dimensions[get_column_letter(i)].width=w
        au.freeze_panes='A2';au.auto_filter.ref=f'A1:J{max(au.max_row,1)}'
        if 'RASTREABILIDADE' in wb.sheetnames:del wb['RASTREABILIDADE']
        tr=wb.create_sheet('RASTREABILIDADE'); tr.append(['Campo','Valor'])
        tr.append(['ID Processamento',process_code or ''])
        tr.append(['Arquivo de origem',Path(source).name])
        tr.append(['Data/hora do processamento',processed_at or datetime.now().isoformat(timespec='seconds')])
        tr.append(['Aba processada',analysis['sheet']])
        tr.append(['Linha do cabeçalho original',analysis['header_idx']+1])
        tr.append(['Registros analisados',analysis['summary']['rows']])
        tr.append(['Ocorrências detectadas',len(analysis['findings'])])
        tr.append(['Registros removidos',len(remove)])
        tr.append(['Observação','Use ID Processamento + ID Hub para rastrear uma decisão entre Excel e SQLite.'])
        for c in tr[1]:c.fill=fill;c.font=font
        tr.column_dimensions['A'].width=34;tr.column_dimensions['B'].width=75
        output.parent.mkdir(parents=True,exist_ok=True);wb.save(output)
        return {'removed':len(remove),'findings':len(analysis['findings'])}
    def _combine(self,d,t):
        try:
            if pd.isna(d) or pd.isna(t):return None
            dd=pd.Timestamp(d).date();tt=pd.to_datetime(t).time() if isinstance(t,str) else (t.time() if isinstance(t,(pd.Timestamp,datetime)) else t)
            return datetime.combine(dd,tt)
        except:return None
    def _key(self,v):return None if pd.isna(v) else (v.isoformat() if isinstance(v,(pd.Timestamp,datetime)) else str(v).strip().lower())
    def _json(self,v):
        if pd.isna(v):return ''
        if isinstance(v,(pd.Timestamp,datetime)):return v.isoformat(sep=' ')
        if isinstance(v,(date,time)):return v.isoformat()
        if hasattr(v,'item'):
            try:return v.item()
            except:pass
        return v
    def _excel(self,v):
        if pd.isna(v):return None
        if isinstance(v,pd.Timestamp):return v.to_pydatetime()
        if hasattr(v,'item'):
            try:return v.item()
            except:pass
        return v
